# -*- coding: utf-8 -*-
"""Ensambla un MANUSCRITO CONSOLIDADO (Caps I-VIII) en un solo .docx, SIN modificar los
archivos individuales (solo los lee). Usa docxcompose para fusionar los .docx con fidelidad
(tablas, ecuaciones, imagenes) y un mini-conversor md->docx para el anexo de borradores que
aun no estan incorporados a los capitulos.
Salida: 11 Redaccion/ICR - Manuscrito consolidado (Caps I-VIII).docx
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxcompose.composer import Composer

BASE = r"C:\Users\Jorge\OneDrive\Escritorio\Claude CODE\ICR"
R = os.path.join(BASE, "11 Redaccion")
CAPS_I_IV = os.path.join(BASE, "00 Proyecto", "Documentos Originales", "Caps I-IV Cortes Morales.docx")

CHAPTERS = [
    CAPS_I_IV,
    os.path.join(R, "Cap V - Calculo de encadenamientos MIP (metodo y resultados).docx"),
    os.path.join(R, "Cap VI - Caracterizacion de los mercados (tipologia descriptiva).docx"),
    os.path.join(R, "Cap VII - Impacto Reforma 2023.docx"),
    os.path.join(R, "Cap VIII - Sintesis y Conclusiones.docx"),
]
JUSTIF = os.path.join(R, "Justificacion - peso del bloque (inserto Cap I) 2026-09-09.docx")
ANEXO_MD = [
    ("A.2  Cap. II — §II.2.5 Criticidad de los productos (borrador)",
     os.path.join(R, "Cap II - subseccion criticidad de productos (borrador para integrar).md")),
    ("A.3  Cap. VII — Marco institucional de México (borrador descriptivo)",
     os.path.join(BASE, "07 Impacto Reforma 2023", "Marco institucional de Mexico y su impacto en las cadenas de valor (descriptivo).md")),
    ("A.4  Cap. VII — Comparativa institucional internacional (borrador descriptivo)",
     os.path.join(BASE, "07 Impacto Reforma 2023", "Comparativa institucional internacional (China y modelo nordico) descriptivo.md")),
]
OUT = os.path.join(R, "ICR - Manuscrito consolidado (Caps I-VIII).docx")

# ---------- helpers de formato ----------
def page_break(doc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); return p

def heading(doc, text, size, before=12, after=6, color=(0x1A,0x1A,0x1A), align=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
    r = p.add_run(text); r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(*color)
    return p

def add_runs(p, text):
    text = re.sub(r'\[\[([^\]|]+)\|([^\]]+)\]\]', r'\2', text)   # [[link|alias]] -> alias
    text = re.sub(r'\[\[([^\]]+)\]\]', r'\1', text)               # [[link]] -> link
    text = re.sub(r'`([^`]+)`', r'\1', text)                      # `code` -> code
    text = re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)', r'\1 (\2)', text)  # [txt](url) -> txt (url)
    pos = 0
    for m in re.finditer(r'\*\*(.+?)\*\*|\*(.+?)\*', text):
        if m.start() > pos: p.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            r = p.add_run(m.group(1)); r.bold = True
        else:
            r = p.add_run(m.group(2)); r.italic = True
        pos = m.end()
    if pos < len(text): p.add_run(text[pos:])

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def add_table(doc, rows):
    headers = [c.strip() for c in rows[0]]
    tbl = doc.add_table(rows=1, cols=len(headers)); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = tbl.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(); add_runs(c.paragraphs[0], h)
        for r in c.paragraphs[0].runs: r.font.bold = True; r.font.size = Pt(8.5)
        shade(c, "E8E8E8")
    for row in rows[1:]:
        cells = tbl.add_row().cells
        for i in range(len(headers)):
            val = row[i].strip() if i < len(row) else ""
            cells[i].text = ""; add_runs(cells[i].paragraphs[0], val)
            for r in cells[i].paragraphs[0].runs: r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def md_to_docx(doc, mdpath):
    lines = open(mdpath, encoding="utf-8").read().split("\n")
    # quitar frontmatter YAML
    if lines and lines[0].strip() == "---":
        end = next((i for i in range(1, len(lines)) if lines[i].strip() == "---"), 0)
        lines = lines[end+1:]
    i = 0; quote_buf = []
    def flush_quote():
        nonlocal quote_buf
        if quote_buf:
            txt = " ".join(quote_buf).strip()
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(14); p.paragraph_format.space_after = Pt(6)
            add_runs(p, txt)
            for r in p.runs: r.font.italic = True; r.font.size = Pt(10.5)
            quote_buf = []
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if s.startswith(">"):
            q = s.lstrip(">").strip()
            q = re.sub(r'^\[![a-zA-Z]+\][+-]?\s*', '', q)   # quitar marcador de callout
            if q: quote_buf.append(q)
            i += 1; continue
        flush_quote()
        if not s:
            i += 1; continue
        if s.startswith("|") and i+1 < len(lines) and re.match(r'^\s*\|?[\s:|-]+\|?\s*$', lines[i+1]):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                block.append(cells); i += 1
            block = [b for b in block if not re.match(r'^[\s:|-]+$', "|".join(b))]  # quitar separador ---
            if block: add_table(doc, block)
            continue
        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#")); txt = s.lstrip("#").strip()
            size = {1:15, 2:13, 3:12, 4:11}.get(lvl, 11)
            heading(doc, txt, size, before=10, after=4, color=(0x33,0x33,0x33))
            i += 1; continue
        if re.match(r'^[-*]\s+', s) and not s.startswith("---"):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, re.sub(r'^[-*]\s+', '', s)); i += 1; continue
        if set(s) <= set("-—* "):   # regla horizontal
            i += 1; continue
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; add_runs(p, s); i += 1
    flush_quote()

# ---------- portada ----------
def cover():
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    sec = doc.sections[0]; sec.page_width = Twips(12240); sec.page_height = Twips(15840)
    for m in ("top_margin","bottom_margin","left_margin","right_margin"): setattr(sec, m, Twips(1701))
    for _ in range(3): doc.add_paragraph()
    heading(doc, "Universidad Autónoma Metropolitana — Unidad Azcapotzalco", 13, align=WD_ALIGN_PARAGRAPH.CENTER)
    heading(doc, "Maestría en Economía", 12, before=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(2): doc.add_paragraph()
    heading(doc, "Los mercados de los minerales críticos en México, 1992–2025:",
            18, before=6, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    heading(doc, "estructura extractiva, cadenas de valor y bases para una política industrial",
            15, before=0, after=6, color=(0x33,0x33,0x33), align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(2): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Idónea Comunicación de Resultados — Manuscrito consolidado (Capítulos I–VIII)").italic = True
    for _ in range(6): doc.add_paragraph()
    for txt in ["Jorge Cortés Morales", "Asesor: Dr. Jordy Micheli Thirion", "Septiembre de 2026"]:
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.add_run(txt)
    for _ in range(2): doc.add_paragraph()
    note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rn = note.add_run("Documento de trabajo generado por consolidación de los capítulos individuales; "
                      "no sustituye a los archivos por capítulo. El Anexo reúne borradores aún no incorporados al cuerpo.")
    rn.italic = True; rn.font.size = Pt(9); rn.font.color.rgb = RGBColor(0x66,0x66,0x66)
    return doc

# ---------- ensamblado ----------
CONTENIDO = [
    "Capítulo I. Introducción",
    "Capítulo II. Marco conceptual y teórico",
    "Capítulo III. La minería mexicana desde la reforma estructural: contexto histórico 1988–2025",
    "Capítulo IV. Estructura empresarial de los diez minerales críticos, 1993–2025",
    "Capítulo V. Diagnóstico de encadenamientos productivos de los minerales críticos",
    "Capítulo VI. Caracterización de los mercados de los minerales críticos",
    "Capítulo VII. La reforma a la Ley Minera de 2023 como contexto institucional, y la referencia internacional",
    "Capítulo VIII. Síntesis, conclusiones y recomendaciones",
    "Anexo. Borradores por integrar (Cap. I, II y VII)",
]
def contenido(doc):
    page_break(doc)
    heading(doc, "Contenido", 15, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    for it in CONTENIDO:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.left_indent = Pt(14)
        add_runs(p, it)

def main():
    master = cover()
    contenido(master)
    comp = Composer(master)
    for ch in CHAPTERS:
        if not os.path.exists(ch): print("  FALTA:", ch, file=sys.stderr); continue
        page_break(master)
        comp.append(Document(ch))
        print("  + capítulo:", os.path.basename(ch))
    # Anexo
    page_break(master)
    heading(master, "Anexo — Borradores por integrar (aún no incorporados a los capítulos)", 15,
            align=WD_ALIGN_PARAGRAPH.CENTER)
    para = master.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rr = para.add_run("Material redactado con fines descriptivos, pendiente de integración por el autor a los "
                      "capítulos correspondientes (Cap. I, II y VII).")
    rr.italic = True; rr.font.size = Pt(10)
    # A.1 inserto de justificación (docx)
    page_break(master)
    heading(master, "A.1  Cap. I — Inserto de justificación: peso del bloque de 10 minerales", 13,
            color=(0x33,0x33,0x33))
    if os.path.exists(JUSTIF):
        comp.append(Document(JUSTIF)); print("  + anexo A.1 (justificación)")
    # A.2-A.4 notas markdown
    for title, mdpath in ANEXO_MD:
        page_break(master)
        heading(master, title, 13, color=(0x33,0x33,0x33))
        if os.path.exists(mdpath):
            md_to_docx(master, mdpath); print("  + anexo:", title.split("—")[0].strip())
        else:
            print("  FALTA md:", mdpath, file=sys.stderr)
    comp.save(OUT)
    print("\nGuardado:", OUT)

if __name__ == "__main__":
    main()
