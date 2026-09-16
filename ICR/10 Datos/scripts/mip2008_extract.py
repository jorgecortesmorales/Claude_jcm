# -*- coding: utf-8 -*-
"""
Extraccion MIP INEGI 2008 (base 2008 / SCIAN 2007) -> CSV cuadrados por CLASE.
Corte de REFERENCIA historica (NO comparable/encadenado con 2013/2018).
Nivel Clase (_4), producto x producto, base DOMESTICA.

Lee los tabulados .XLSX (con filas de titulo y encabezados en filas/columnas
propias, distinto del CSV abierto de 2013/2018) ubicando DINAMICAMENTE los
encabezados, y vuelca a CSV con el formato reutilizable por mip2008_calc.py:
  - Z_dom.csv    flujos domesticos producto x producto (mipdpb_pxp_4)
  - A_inegi.csv  coef. tecnicos domesticos de INEGI    (mipdctcpxp_4)  [validacion]
  - L_inegi.csv  (I-A)^-1 de INEGI                      (mipdcdipxp_4)  [validacion]
  - vec.csv      por clase: vbp (utilizacion total prod. interna) y DI domestica
  - codes.csv    orden, codigo, nombre
Herramientas: openpyxl 3.1.5 + numpy 2.2.5 (sin Node, sin pandas).
"""
import csv, os, re
import openpyxl

BASE = r"C:\Users\Jorge\OneDrive\Escritorio\Claude CODE\ICR\10 Datos"
SRC  = os.path.join(BASE, "Bases Originales", "10 MIP INEGI",
                    "2003_2008 historicas (no comparables)",
                    "_extract_2008", "Tabulados_mip2008", "mip2008")
OUT  = os.path.join(BASE, "processed", "mip2008_intermedios")

FILES = {
    "Z_dom": "mipdpb_pxp_4.XLSX",   # flujos domesticos (millones de pesos)
    "A":     "mipdctcpxp_4.XLSX",   # coeficientes tecnicos domesticos (A)
    "L":     "mipdcdipxp_4.XLSX",   # directos+indirectos = (I-A)^-1 (L)
    "T":     "miptpb_pxp_4.XLSX",   # flujos totales (para VBP = utilizacion total)
}

CODE_RE = re.compile(r"^\d{6}$")
INT_RE  = re.compile(r"^\d+$")

def load_grid(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    grid = [[(c if c is not None else "") for c in row]
            for row in ws.iter_rows(values_only=True)]
    wb.close()
    return grid

def s(v):
    return str(v).strip()

def find_col_code_row(grid):
    """Fila (0-based) con mas codigos de 6 digitos en columnas >=3 dentro de las primeras 15."""
    best = (-1, -1)
    for i in range(min(15, len(grid))):
        cnt = sum(1 for c in grid[i][3:] if CODE_RE.match(s(c)))
        if cnt > best[0]:
            best = (cnt, i)
    return best[1]

def parse_square(path):
    """Devuelve (codes_orden, name_by_code, M) del bloque producto x producto.
    Filas de datos: col0 entero (NO.) y col1 codigo 6-dig. Columnas: fila de codigos."""
    grid = load_grid(path)
    crow = find_col_code_row(grid)
    col_of_code = {}
    for j, c in enumerate(grid[crow]):
        if j >= 3 and CODE_RE.match(s(c)):
            cc = s(c)
            col_of_code.setdefault(cc, j)
    codes, name_by_code, row_vals = [], {}, {}
    for r in grid:
        if INT_RE.match(s(r[0])) and CODE_RE.match(s(r[1])):
            code = s(r[1])
            if code in row_vals:
                continue
            codes.append(code)
            name_by_code[code] = s(r[2])
            row_vals[code] = r
    assert set(codes) == set(col_of_code), \
        f"desajuste fila/col en {os.path.basename(path)}: {len(codes)} vs {len(col_of_code)}"
    n = len(codes)
    M = [[0.0] * n for _ in range(n)]
    for i, ci in enumerate(codes):
        r = row_vals[ci]
        for jj, cj in enumerate(codes):
            v = s(r[col_of_code[cj]])
            M[i][jj] = float(v) if v not in ("", "-") else 0.0
    return codes, name_by_code, M

def parse_vectors(path):
    """De un archivo de flujos (pb/total): x = 'UTILIZACION TOTAL ... PRODUCCION INTERNA';
    di = 'TOTAL DEMANDA INTERMEDIA'. Localiza las columnas por texto de encabezado."""
    grid = load_grid(path)
    # fila de nombres de columna = una arriba de la fila de codigos
    crow = find_col_code_row(grid)
    hdr = grid[crow - 1]
    def find_col(substr):
        for j, c in enumerate(hdr):
            if substr in s(c).upper():
                return j
        return None
    jx = find_col("UTILIZACI")           # utilizacion total de la produccion interna
    jdi = find_col("TOTAL DEMANDA INTERMEDIA")
    x, di = {}, {}
    for r in grid:
        if INT_RE.match(s(r[0])) and CODE_RE.match(s(r[1])):
            code = s(r[1])
            def num(j):
                if j is None:
                    return 0.0
                v = s(r[j]); return float(v) if v not in ("", "-") else 0.0
            x.setdefault(code, num(jx))
            di.setdefault(code, num(jdi))
    return x, di

def write_matrix(path, codes, M):
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["codigo"] + codes)
        for i, ci in enumerate(codes):
            w.writerow([ci] + [repr(v) for v in M[i]])  # repr => round-trip exacto

def main():
    os.makedirs(OUT, exist_ok=True)
    print("Extrayendo MIP 2008 (nivel Clase, domestica)...")

    codes, name, Z = parse_square(os.path.join(SRC, FILES["Z_dom"]))
    cA, _, A = parse_square(os.path.join(SRC, FILES["A"]))
    cL, _, L = parse_square(os.path.join(SRC, FILES["L"]))
    assert codes == cA == cL, "orden/base de clases distinto entre Z, A y L"
    n = len(codes)
    print(f"  n clases = {n}  (2013=822, 2018=834 como referencia)")

    # VBP y DI domestica: x desde tabla TOTAL (utilizacion total prod. interna),
    # di domestica desde tabla de flujos DOMESTICA.
    xT, _   = parse_vectors(os.path.join(SRC, FILES["T"]))
    xD, diD = parse_vectors(os.path.join(SRC, FILES["Z_dom"]))
    # coherencia: la 'utilizacion total de la produccion interna' debe coincidir
    maxdif = max(abs(xT.get(c, 0.0) - xD.get(c, 0.0)) for c in codes)
    print(f"  max|VBP_total - VBP_dom| = {maxdif:.3e}  (deben coincidir; VBP = prod. interna)")
    x = xT

    write_matrix(os.path.join(OUT, "Z_dom.csv"), codes, Z)
    write_matrix(os.path.join(OUT, "A_inegi.csv"), codes, A)
    write_matrix(os.path.join(OUT, "L_inegi.csv"), codes, L)
    with open(os.path.join(OUT, "vec.csv"), "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh); w.writerow(["codigo", "nombre", "vbp", "di_domestica"])
        for c in codes:
            w.writerow([c, name[c], repr(x[c]), repr(diD[c])])
    with open(os.path.join(OUT, "codes.csv"), "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh); w.writerow(["orden", "codigo", "nombre"])
        for i, c in enumerate(codes):
            w.writerow([i, c, name[c]])
    print(f"  escritos en {OUT}")
    print("  archivos: Z_dom.csv A_inegi.csv L_inegi.csv vec.csv codes.csv")

if __name__ == "__main__":
    main()
