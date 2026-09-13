# -*- coding: utf-8 -*-
"""Ensambla el MANUSCRITO CONSOLIDADO usando el motor de Word (COM), que produce un .docx
que Word abre sin errores (docxcompose generaba archivos que Word marcaba como corruptos).
Reutiliza los helpers de build_manuscrito.py para las piezas propias (portada, contenido,
anexo md->docx) y usa Word Selection.InsertFile para fusionar con fidelidad total.
NO modifica los capitulos individuales (solo los inserta). Salidas: .docx y .pdf.
"""
import os, sys, tempfile, glob
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_manuscrito as bm
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client as win32

R = bm.R
OUT_DOCX = os.path.join(R, "ICR - Manuscrito consolidado (Caps I-VIII).docx")
OUT_PDF  = os.path.join(R, "ICR - Manuscrito consolidado (Caps I-VIII).pdf")

def build_piece(fn):
    """crea un docx temporal con python-docx usando una funcion fn(doc)"""
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    fn(doc)
    fd, path = tempfile.mkstemp(suffix=".docx", prefix="_ms_"); os.close(fd)
    doc.save(path); return path

def cover_piece(doc):
    # reutiliza la portada + contenido de build_manuscrito
    for _ in range(3): doc.add_paragraph()
    bm.heading(doc, "Universidad Autónoma Metropolitana — Unidad Azcapotzalco", 13, align=WD_ALIGN_PARAGRAPH.CENTER)
    bm.heading(doc, "Maestría en Economía", 12, before=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(2): doc.add_paragraph()
    bm.heading(doc, "Los mercados de los minerales críticos en México, 1992–2025:", 18, before=6, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    bm.heading(doc, "estructura extractiva, cadenas de valor y bases para una política industrial", 15, before=0, after=6, color=(0x33,0x33,0x33), align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(2): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Idónea Comunicación de Resultados — Manuscrito consolidado (Capítulos I–VIII)").italic = True
    for _ in range(6): doc.add_paragraph()
    for txt in ["Jorge Cortés Morales", "Asesor: Dr. Jordy Micheli Thirion", "Septiembre de 2026"]:
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.add_run(txt)
    for _ in range(2): doc.add_paragraph()
    note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rn = note.add_run("Documento de trabajo generado por consolidación de los capítulos individuales; "
                      "no sustituye a los archivos por capítulo. El Anexo reúne borradores aún no incorporados al cuerpo.")
    rn.italic = True; rn.font.size = Pt(9); rn.font.color.rgb = RGBColor(0x66,0x66,0x66)
    # Contenido
    bm.page_break(doc)
    bm.heading(doc, "Contenido", 15, align=WD_ALIGN_PARAGRAPH.CENTER); doc.add_paragraph()
    for it in bm.CONTENIDO:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(6); pp.paragraph_format.left_indent = Pt(14)
        bm.add_runs(pp, it)

def annex_pre_piece(doc):
    bm.heading(doc, "Anexo — Borradores por integrar (aún no incorporados a los capítulos)", 15, align=WD_ALIGN_PARAGRAPH.CENTER)
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rr = para.add_run("Material redactado con fines descriptivos, pendiente de integración por el autor a los capítulos correspondientes (Cap. I, II y VII).")
    rr.italic = True; rr.font.size = Pt(10)
    bm.page_break(doc)
    bm.heading(doc, "A.1  Cap. I — Inserto de justificación: peso del bloque de 10 minerales", 13, color=(0x33,0x33,0x33))

def annex_md_piece(doc):
    for title, mdpath in bm.ANEXO_MD:
        bm.page_break(doc)
        bm.heading(doc, title, 13, color=(0x33,0x33,0x33))
        if os.path.exists(mdpath): bm.md_to_docx(doc, mdpath)

def main():
    tmp = []
    cover = build_piece(cover_piece); tmp.append(cover)
    annex_pre = build_piece(annex_pre_piece); tmp.append(annex_pre)
    annex_md = build_piece(annex_md_piece); tmp.append(annex_md)
    # secuencia de inserción
    seq = [cover] + bm.CHAPTERS + [annex_pre, bm.JUSTIF, annex_md]
    seq = [p for p in seq if os.path.exists(p)]

    word = win32.Dispatch("Word.Application"); word.Visible = False; word.DisplayAlerts = 0
    try:
        doc = word.Documents.Add()
        sel = word.Selection
        for k, path in enumerate(seq):
            sel.EndKey(6)                 # wdStory: ir al final
            if k > 0: sel.InsertBreak(7)  # wdPageBreak
            sel.InsertFile(FileName=os.path.abspath(path), ConfirmConversions=False, Link=False)
            print("  + insertado:", os.path.basename(path))
        doc.SaveAs(os.path.abspath(OUT_DOCX), FileFormat=16)  # wdFormatDocumentDefault (.docx)
        doc.SaveAs(os.path.abspath(OUT_PDF),  FileFormat=17)  # wdFormatPDF
        pags = doc.ComputeStatistics(2)                       # wdStatisticPages
        doc.Close(False)
        print(f"\nGuardado DOCX y PDF. Páginas: {pags}")
    finally:
        word.Quit()
    for p in tmp:
        try: os.remove(p)
        except: pass

if __name__ == "__main__":
    main()
