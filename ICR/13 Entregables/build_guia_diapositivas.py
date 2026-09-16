# -*- coding: utf-8 -*-
"""Guía de exposición a fondo — diapositivas 3, 7, 8, 10 y 16 del protocolo ICR.
Combina el contenido del proyecto con el marco conceptual amplio (autores, mecanismos,
supuestos y críticas) necesario para entender y defender cada tema. python-docx, formato
tesis (Times New Roman 12, Carta, márgenes 3 cm)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTDIR = r"C:\Users\Jorge\OneDrive\Escritorio\Claude CODE\ICR\13 Entregables\Resumenes descriptivos"
FECHA = "2026-09-07"
OUT = os.path.join(OUTDIR, f"Guia de exposicion - diapositivas 3,7,8,10,16 {FECHA}.docx")
INK = RGBColor(0x21,0x1d,0x18); COP = RGBColor(0xb0,0x57,0x1e); MUT = RGBColor(0x6b,0x64,0x5a)

doc = Document()
st = doc.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.15
sec=doc.sections[0]; sec.page_width=Cm(21.59); sec.page_height=Cm(27.94)
for m in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(sec,m,Cm(3))
for lvl in (1,2,3):
    h=doc.styles['Heading %d'%lvl]; h.font.name='Times New Roman'; h.font.color.rgb=INK

def P(text=None, size=12, bold=False, italic=False, color=None, first=False, after=6):
    p=doc.add_paragraph(); pf=p.paragraph_format; pf.space_after=Pt(after); pf.line_spacing=1.15
    if first: pf.first_line_indent=Pt(21)
    if text is not None:
        r=p.add_run(text); r.font.size=Pt(size); r.bold=bold; r.italic=italic
        if color is not None: r.font.color.rgb=color
    return p

def runs(parts, first=False, after=6):
    p=doc.add_paragraph(); pf=p.paragraph_format; pf.space_after=Pt(after); pf.line_spacing=1.15
    if first: pf.first_line_indent=Pt(21)
    for t,opt in parts:
        r=p.add_run(t); r.font.size=Pt(opt.get('size',12)); r.bold=opt.get('b',False); r.italic=opt.get('i',False)
        if opt.get('color') is not None: r.font.color.rgb=opt['color']
    return p

def H1(t):
    p=doc.add_heading(level=1); r=p.add_run(t); r.font.size=Pt(15); r.bold=True; r.font.color.rgb=INK
def H2(t):
    p=doc.add_heading(level=2); r=p.add_run(t); r.font.size=Pt(13); r.bold=True; r.font.color.rgb=COP
def H3(t):
    p=doc.add_heading(level=3); r=p.add_run(t); r.font.size=Pt(12); r.bold=True; r.font.color.rgb=INK
def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); r=p.add_run(text); r.font.size=Pt(11.5); r.font.name='Times New Roman'
def qa(q,a):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    r=p.add_run('P: '); r.bold=True; r.font.size=Pt(11.5); r.font.name='Times New Roman'; r.font.color.rgb=COP
    r2=p.add_run(q); r2.font.size=Pt(11.5); r2.italic=True; r2.font.name='Times New Roman'
    p2=doc.add_paragraph(); p2.paragraph_format.space_after=Pt(8); p2.paragraph_format.left_indent=Cm(0.5)
    r3=p2.add_run('R: '); r3.bold=True; r3.font.size=Pt(11.5); r3.font.name='Times New Roman'
    r4=p2.add_run(a); r4.font.size=Pt(11.5); r4.font.name='Times New Roman'

def guion(text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(0.5); p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.space_before=Pt(2)
    r=p.add_run('« '+text+' »'); r.italic=True; r.font.size=Pt(11.5); r.font.name='Times New Roman'; r.font.color.rgb=MUT

# ============ PORTADA ============
P('ICR · Maestría en Economía · UAM Azcapotzalco', size=10.5, color=COP, after=2)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
r=p.add_run('Guía de exposición a fondo'); r.bold=True; r.font.size=Pt(20); r.font.name='Times New Roman'; r.font.color.rgb=INK
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(10)
r=p.add_run('Diapositivas 3, 7, 8, 10 y 16 del protocolo — todo lo necesario para entenderlas y defenderlas'); r.font.size=Pt(13.5); r.italic=True; r.font.color.rgb=MUT
pr=doc.add_paragraph(); pPr=pr._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); bot=OxmlElement('w:bottom')
bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'12'); bot.set(qn('w:space'),'6'); bot.set(qn('w:color'),'b0571e')
pbdr.append(bot); pPr.append(pbdr); pr.paragraph_format.space_after=Pt(10)
P('Esta guía va más allá de lo que dicen las diapositivas: explica el marco conceptual de fondo —los autores, los mecanismos, los supuestos y las críticas— que hace falta para comprender cada tema y responder con solvencia. Cada diapositiva se aborda en seis pasos: qué muestra, el marco a fondo, cómo se conecta con la tesis, un guion breve para exponer, las preguntas probables del jurado con su respuesta, y los errores finos a evitar. Las cifras y bases se detallan en el «Resumen metodológico y de resultados»; aquí el foco es entender y explicar.', first=True)

# ==================================================================
# DIAPOSITIVA 3 — JUSTIFICACIÓN / BRECHA
# ==================================================================
H1('Diapositiva 3 — Justificación: tres literaturas que no se cruzan')

H2('Qué muestra')
P('Que el tema se ha estudiado desde tres frentes que casi no dialogan entre sí, y que la tesis tiende un puente construyendo datos e indicadores por mineral. Cierra con el ancla teórica (Prebisch) y la pertinencia al campo EFI.')

H2('El marco a fondo')
H3('Las tres literaturas y qué le falta a cada una')
runs([('1) Economía política de la privatización y liberalización (1988-1993). ',{'b':True}),
 ('Documenta cómo se desincorporaron las reservas y se privatizaron las paraestatales, y la concentración empresarial resultante. Es histórico-institucional: explica el “cómo llegamos aquí”, pero no mide la cadena de valor ni la estructura de mercado con indicadores. Autores: Sariego (2009); Delgado Wise y Del Pozo (2001).',{})])
runs([('2) Estudios ambientales y sociales del extractivismo. ',{'b':True}),
 ('Analizan el impacto territorial, social y ecológico de la gran minería y el modelo primario-exportador. Son críticos y ricos en lo cualitativo, pero con poca medición económica desagregada. Autor de referencia en México: Azamar y Ponce (2014); en la región, la corriente del neo-extractivismo.',{})])
runs([('3) Estudios macroeconómicos del sector. ',{'b':True}),
 ('Cuantifican el aporte de la minería al PIB, al empleo y a las exportaciones (agregados nacionales). Miden el “cuánto” global, pero no bajan a la estructura ni a la cadena por mineral. Autores: Aliphat Rodríguez et al. (2025); reportes de CAMIMEX y USGS.',{})])
P('La brecha, entonces, no es que falten estudios, sino que ninguno construye —de forma sistemática y desagregada por mineral— los datos e indicadores que unen estructura de mercado + cadena de valor + inserción global. Ese acervo empírico es el aporte de la tesis.')

H3('El telón de fondo teórico: estructuralismo y la tesis Prebisch-Singer')
P('La relevancia teórica se ancla en el estructuralismo latinoamericano (CEPAL), cuyo referente es Raúl Prebisch (1950). Su idea central, la tesis Prebisch-Singer, sostiene que los términos de intercambio de los exportadores de materias primas se deterioran de forma secular frente a las manufacturas. Dos mecanismos la explican:')
bullet('Por el lado de la demanda: la elasticidad-ingreso de las materias primas es baja (ley de Engel) y la de las manufacturas es alta; al crecer el ingreso mundial, la demanda se desplaza hacia las manufacturas y los precios relativos de los primarios caen.')
bullet('Por el lado de la distribución: las ganancias de productividad se retienen en el “centro” (como salarios y beneficios, por poder sindical y oligopólico), mientras la “periferia” traslada las suyas a menores precios de exportación.')
P('La implicación clásica fue la industrialización (sustitución de importaciones). Para esta tesis, Prebisch-Singer aporta el marco macro del problema —exportar barato lo primario y comprar caro lo industrial—; el enclave estructural (diapositiva 7) lo lleva al plano micro-sectorial. Conviene conocer la crítica: el debate empírico sobre si el deterioro es universal o depende del producto y del periodo (por ejemplo, Grilli y Yang), y la alta volatilidad de los precios de commodities.')

H3('Pertinencia al campo EFI')
P('El campo es Empresas, Finanzas e Innovación. La tesis caracteriza empresas y grupos del sector, su estructura de industria, su integración vertical y la organización de sus cadenas de valor: todos objetos centrales de la economía industrial, que es el corazón de EFI.')

H2('Cómo exponerla')
guion('El tema se ha visto desde tres frentes que no se hablan: la economía política de la privatización, los estudios ambientales del extractivismo y los macroeconómicos del aporte al PIB. Ninguno construye, mineral por mineral, los indicadores de estructura y de cadena de valor. Eso es lo que aporto. Todo anclado en la tradición estructuralista de Prebisch y en el campo EFI.')

H2('Preguntas probables del jurado')
qa('¿Por qué Prebisch y no directamente la teoría de la dependencia o del enclave?',
   'Prebisch-Singer da el marco macro (deterioro de términos de intercambio del exportador primario). El enclave estructural, en la diapositiva 7, es la forma concreta y sectorial de ese problema en México. Son complementarios: uno enmarca, el otro operacionaliza.')
qa('¿No existen ya análisis insumo-producto de la minería mexicana?',
   'A nivel agregado sí, pero no por mineral y a nivel de Clase SCIAN, ni cruzados con concentración, captura de valor y comercio por etapa. El aporte es esa desagregación y ese cruce sistemático.')
qa('¿Qué agrega frente a los reportes de CAMIMEX o del USGS?',
   'Esos reportan producción, precios y estructura básica. Aquí se construyen indicadores de encadenamiento (Leontief, Ghosh), de captura de valor (CCV) y de inserción por etapa, que esos reportes no calculan.')

H2('Errores a evitar')
bullet('No presentar las tres literaturas como “malas”: son valiosas pero parciales; el punto es que no se cruzan.')
bullet('No prometer explicación causal: el aporte es descriptivo y empírico (construir y medir), como se aclara en la diapositiva 8.')

# ==================================================================
# DIAPOSITIVA 7 — MARCO TEÓRICO
# ==================================================================
H1('Diapositiva 7 — Marco teórico')
P('Existen dos versiones de esta diapositiva. Domina ambas: la del protocolo (estructuralismo + organización industrial y cadenas de valor) y la del estado actual de la tesis (el enclave estructural). Se complementan.', italic=True, color=MUT)

H2('Qué muestra')
P('En la versión del protocolo: el ancla estructuralista (Prebisch) más cuatro marcos analíticos que operan como lentes descriptivos. En la versión actual: el paso del enclave clásico (Cardoso-Faletto) al enclave estructural (con respaldo neo-extractivista de Svampa) como concepto ordenador.')

H2('El marco a fondo')
H3('Los cuatro marcos analíticos (versión del protocolo)')
runs([('Economía industrial y de la organización — el paradigma Estructura-Conducta-Desempeño (Mason, 1939; Bain, 1956). ',{'b':True}),
 ('Sostiene que la estructura del mercado (número y tamaño de empresas, barreras a la entrada) condiciona la conducta de las firmas y, con ella, el desempeño (precios, márgenes, eficiencia). Bain formalizó las barreras a la entrada. En la tesis, esto justifica medir la estructura con el HHI: la concentración es el punto de partida para leer el comportamiento del mercado.',{})])
runs([('Integración vertical y costos de transacción (Coase, 1937; Williamson, 1985; Klein, Crawford y Alchian, 1978). ',{'b':True}),
 ('Explica por qué una empresa decide hacer internamente una etapa de la cadena en vez de comprarla en el mercado. Cuando hay activos específicos y “cuasi-rentas apropiables”, el intercambio en el mercado se vuelve riesgoso (problema del hold-up), y conviene integrar verticalmente. Esto ilumina por qué los grupos mineros integran hasta la fundición y refinación (la mina y la fundición son activos muy específicos entre sí) pero no necesariamente más allá.',{})])
runs([('Economía institucional (North, 1990). ',{'b':True}),
 ('Las instituciones son “las reglas del juego”: normas formales (leyes, concesiones) e informales que estructuran los incentivos y reducen la incertidumbre. La Ley Minera de 1992 y su reforma de 2023 son instituciones formales que moldean el modelo extractivo. El cambio institucional es dependiente de la trayectoria (path dependence).',{})])
runs([('Cadenas de valor globales (Gereffi, 1994; Kaplinsky y Morris, 2001; posición: Antràs y Chor, 2013). ',{'b':True}),
 ('Estudia cómo se organiza y gobierna la producción fragmentada a escala mundial. Gereffi distingue cadenas “dirigidas por el comprador” y “por el productor”, y una tipología de gobernanza (mercado, modular, relacional, cautiva, jerárquica). Antràs y Chor formalizan la posición en la cadena (upstreamness / downstreamness): qué tan “aguas arriba” o “aguas abajo” está una actividad respecto del consumo final. México, exportando mena/concentrado, está muy aguas arriba.',{})])
P('Clave para la defensa: estos marcos, aunque en su origen sirven para análisis causales, aquí se usan en clave descriptiva —como lentes para describir la estructura y la cadena—, no para estimar relaciones causales.')

H3('El enclave: del clásico al estructural (versión actual)')
runs([('Enclave clásico (Cardoso y Faletto, 1969). ',{'b':True}),
 ('En la teoría de la dependencia, la “economía de enclave” es un sector exportador controlado por capital extranjero, débilmente vinculado al resto de la economía nacional, cuyas ganancias se repatrían. El rasgo definitorio es la propiedad extranjera: la renta sale del país con la empresa.',{})])
runs([('Enclave estructural (concepto ordenador de la tesis). ',{'b':True}),
 ('En México el capital es mayoritariamente nacional (Grupo México, Peñoles, Autlán), y sin embargo la cadena se desconecta aguas abajo: se extrae y se exporta en bruto sin desarrollar la transformación industrial local. El “enclave” ya no se define por quién es el dueño, sino por dónde se detiene la cadena. Es un desplazamiento del criterio: de la propiedad a la estructura del encadenamiento.',{})])
runs([('Respaldo neo-extractivista (Svampa, 2013). ',{'b':True}),
 ('El “Consenso de los Commodities” describe cómo, durante el boom de materias primas, América Latina profundizó su perfil primario-exportador (reprimarización), incluso bajo gobiernos de distinto signo. Da el marco contemporáneo del extractivismo en el que se inscribe el enclave estructural.',{})])
P('El giro que conviene enunciar con claridad: el problema mexicano no es de quién es la mina, sino dónde se detiene la cadena. Y —punto fino— un encadenamiento estadístico alto (Ghosh) no equivale a una cadena desarrollada: por eso el concepto se sostiene en la lectura conjunta de varios indicadores, no en uno solo.')

H2('Cómo exponerla')
guion('Versión actual: parto del enclave clásico de Cardoso y Faletto —definido por la propiedad extranjera— y, por contraste, defino el enclave estructural: capital nacional pero cadena desconectada aguas abajo. Lo respaldo con el neo-extractivismo de Svampa y lo hago medible con HHI, Ghosh y el coeficiente de captura de valor.')
guion('Versión protocolo: lo anclo en el estructuralismo de Prebisch y opero con cuatro lentes: economía industrial para la estructura (HHI), integración vertical para saber hasta dónde integran las firmas, economía institucional para el marco legal, y cadenas de valor globales para la inserción externa.')

H2('Preguntas probables del jurado')
qa('¿El enclave estructural es lo mismo que el enclave de Cardoso y Faletto?',
   'No. El clásico se define por la propiedad extranjera del capital. El estructural mantiene el capital nacional pero constata la desconexión de la cadena aguas abajo. Se llega a él por contraste con el clásico.')
qa('Si el diseño es descriptivo, ¿por qué usar teorías con vocación causal (Estructura-Conducta-Desempeño, costos de transacción)?',
   'Se usan como marcos de referencia para describir y ordenar, no para estimar efectos. Aportan el vocabulario (estructura, integración vertical, gobernanza de cadena) con el que se interpreta lo que muestran los indicadores.')
qa('¿Cómo se relaciona Prebisch con el enclave estructural?',
   'Prebisch-Singer es el marco macro (deterioro de términos de intercambio del exportador primario). El enclave estructural es su expresión micro-sectorial: por qué México se queda en el eslabón primario y no captura el valor aguas abajo.')
qa('¿No es criticable el modelo de Ghosh como base teórica?',
   'Sí, y por eso se usa con cautela: Ghosh es más defendible como modelo de precios (Dietzenbacher, 1997) que como predicción de cantidades. Aquí es un descriptor de posición estructural, cruzado siempre con el CCV, el comercio por etapa y el mapa de empresas.')

H2('Errores a evitar')
bullet('No mezclar las dos versiones sin avisar: si expones el enclave estructural, deja claro que es el concepto ordenador actual; el protocolo original se ancla en Prebisch + economía industrial y CVG.')
bullet('No atribuir a Cardoso-Faletto el “enclave estructural”: ellos definen el clásico (propiedad extranjera); el estructural es la aportación conceptual de la tesis.')
bullet('No presentar el Ghosh alto como prueba de cadena desarrollada.')

# ==================================================================
# DIAPOSITIVA 8 — ENFOQUE Y HORIZONTE
# ==================================================================
H1('Diapositiva 8 — Enfoque y horizonte temporal')

H2('Qué muestra')
P('Que el enfoque es descriptivo y de métodos mixtos; que la unidad de observación es el mineral-año (diez minerales fijos); y que el horizonte es 1992-2025, con un detalle temporal que depende de la fuente.')

H2('El marco a fondo')
H3('Descriptivo vs. causal: por qué descriptivo')
P('Un diseño causal busca estimar el efecto de una variable sobre otra (por ejemplo, si más concentración “causa” menos encadenamiento) e implica una estrategia de identificación creíble (variación exógena, contrafactual). Un diseño descriptivo busca caracterizar con rigor cómo es un fenómeno. Aquí se elige el descriptivo por dos razones: (a) la información encadenable por mineral son apenas dos o tres cortes de la Matriz Insumo-Producto, insuficientes para una identificación causal creíble; y (b) el aporte que la literatura no ha hecho es, precisamente, construir y describir los indicadores. Una relación causal quedaría como agenda futura.')
H3('Métodos mixtos')
P('Combina lo cuantitativo (indicadores construidos desde fuentes documentales: MIP, comercio, precios) con lo cualitativo (lectura de reportes corporativos y, opcionalmente, entrevistas). La lógica de métodos mixtos (Creswell y Clark) es que cada componente cubre lo que el otro no ve: el dato agregado no revela la integración intra-firma, y el reporte de empresa sí. Para lo cualitativo-causal de gabinete se apoya en el process tracing (Beach y Pedersen) y el análisis temático (Braun y Clarke).')
H3('Unidad de observación: el mineral-año')
P('Cada observación es un mineral en un año concreto (por ejemplo, “cobre-2018”). Esta unidad permite dos lecturas: comparar entre minerales en un año y seguir cada mineral a lo largo del tiempo. Es la rejilla común que hace comparables los indicadores.')
H3('El horizonte 1992-2025 y por qué la cobertura es heterogénea')
P('1992 marca la Ley Minera que configuró el modelo actual; 2025 es el cierre de la información. Son 34 años. Pero no todos los indicadores cubren todo el periodo con la misma granularidad: el CCV y el comercio por etapa son series anuales continuas, mientras que los encadenamientos de la MIP son “fotos” de años base concretos. Enunciar esto por adelantado es una fortaleza metodológica, no una debilidad: se trabaja con la granularidad que cada fuente permite y se declara.')

H2('Cómo exponerla')
guion('El enfoque es descriptivo y de métodos mixtos: observo cada mineral año con año, para los diez minerales, de 1992 —la Ley Minera— a 2025. No busco un efecto causal, sino caracterizar con rigor los mercados y su cadena de valor. El detalle temporal depende de la fuente: unos indicadores son series anuales; otros, cortes de la matriz insumo-producto.')

H2('Preguntas probables del jurado')
qa('¿Por qué descriptivo y no un modelo econométrico?',
   'Porque una identificación causal creíble exige más variación de la que hay (dos o tres cortes de MIP por mineral), y porque el vacío en la literatura es descriptivo: no existen estos indicadores por mineral. Un modelo causal es agenda futura, no el aporte de esta tesis.')
qa('¿Por qué la unidad es el mineral-año y no la empresa o el estado?',
   'Porque la pregunta es sobre los mercados de cada mineral y su evolución; el mineral-año permite comparar entre minerales y en el tiempo. El nivel de empresa (cadena local) y de estado (regionalización) entran como análisis complementarios.')
qa('¿Por qué empezar en 1992?',
   'Porque la Ley Minera de 1992 configuró el régimen de concesiones y el modelo extractivo vigente; es el punto de partida institucional del fenómeno que se describe.')

H2('Errores a evitar')
bullet('No afirmar “serie continua 1992-2025 para todos los indicadores”: la MIP son cortes; declararlo (ver diapositiva 16).')
bullet('No confundir “descriptivo” con “superficial”: es descriptivo en el sentido metodológico (no causal), pero con rigor cuantitativo.')

# ==================================================================
# DIAPOSITIVA 10 — TÉCNICAS
# ==================================================================
H1('Diapositiva 10 — Técnicas de análisis: cuatro descriptores')

H2('Qué muestra')
P('Que cada objetivo tiene su técnica: O1 (estructura) con el HHI; O1 (encadenamiento) con Leontief, Ghosh y el CCV; O2 (cadena local) con la demanda intermedia; O3 (inserción global) con el comercio por etapa.')

H2('El marco a fondo')
H3('HHI — concentración (Objetivo 1)')
P('El índice de Herfindahl-Hirschman es la suma de los cuadrados de las participaciones de mercado. Elevar al cuadrado hace que los grandes pesen mucho más que los pequeños, de modo que distingue un mercado con un líder dominante de otro repartido entre iguales. Umbrales habituales (Departamento de Justicia de EE. UU. / Comisión Europea): por debajo de 1 500, baja concentración; entre 1 500 y 2 500, moderada; por encima de 2 500, alta; 10 000 es monopolio. Un atajo útil: el “número efectivo de competidores” es 1 dividido entre el HHI en fracción.')
H3('Leontief y Ghosh — encadenamiento (Objetivo 1)')
P('Ambos se derivan de la Matriz Insumo-Producto, que registra cuánto le vende cada sector a cada otro. El modelo de Leontief (1941) es de demanda: a partir de coeficientes técnicos (cuánto insumo requiere cada sector) mide los encadenamientos hacia atrás —cuánto arrastra un sector a sus proveedores—. El modelo de Ghosh (1958) es su dual, de oferta: a partir de coeficientes de distribución (a dónde va la producción de cada sector) mide los encadenamientos hacia adelante —cuánto empuja un sector a las industrias que lo usan como insumo—. Los índices se normalizan con la fórmula de Hirschman-Rasmussen (media de la economía = 1) para comparar sectores de tamaños distintos. Hirschman (1958) es el origen conceptual de los “eslabonamientos”; Rasmussen (1956), de los índices normalizados.')
P('Dos cautelas teóricas que conviene tener a la mano: el modelo de Ghosh ha sido cuestionado como modelo de cantidades (Oosterhaven, 1988) y se defiende mejor como modelo de precios (Dietzenbacher, 1997); por eso aquí se usa como descriptor de posición, no como predicción. Y un valor alto de Ghosh indica arrastre estructural, no captura de valor: puede ser un solo eslabón (la fundición) y no una cadena larga.')
H3('CCV — captura de valor (Objetivo 1)')
P('El coeficiente de captura de valor es un aporte de la tesis: el cociente entre el valor unitario de lo que México exporta en bruto y el precio del producto de referencia refinado. Da la versión anual y continua del encadenamiento hacia adelante, que el Ghosh —atado a los cortes de la MIP— no ofrece. Es informativo sobre todo en metales base (cobre, zinc).')
H3('Demanda intermedia — cadena local (Objetivo 2)')
P('Con la misma MIP se lee la fila de cada mineral: qué sectores lo compran como insumo dentro del país. Eso responde si existe industria usuaria doméstica. Se complementa con un mapa de empresas de transformación (a nivel de firma), porque una parte de la cadena ocurre dentro de los propios grupos (integración vertical) y el dato sectorial la subestima.')
H3('Comercio por etapa — inserción global (Objetivo 3)')
P('Se clasifica cada fracción arancelaria en su etapa (mena → concentrado → refinado → intermedio → bien final) y se mide la participación de cada etapa en las exportaciones. El análisis espejo compara lo que México exporta con lo que importa del mismo mineral: exportar en bruto e importar el producto procesado es la firma de que la transformación ocurre fuera. Es la lógica de la especialización vertical en el comercio (Hummels, Ishii y Yi), y conecta con las medidas de valor agregado en el comercio (TiVA) que la tesis usa en su avance.')

H2('Cómo exponerla')
guion('Cada objetivo tiene su herramienta. Para la estructura, el HHI. Para el encadenamiento, Leontief hacia atrás y Ghosh hacia adelante desde la matriz insumo-producto, más el coeficiente de captura de valor como serie anual. Para la cadena local, la demanda intermedia y el mapa de empresas. Para la inserción global, el comercio por etapa con análisis espejo. Todos son descriptores, no un modelo causal.')

H2('Preguntas probables del jurado')
qa('¿Por qué el Ghosh (hacia adelante) y no solo Leontief (hacia atrás)?',
   'Porque la pregunta de la tesis es si el mineral alimenta una industria transformadora dentro del país; eso es hacia adelante. Leontief (proveedores) se reporta como complemento, pero en minería es bajo por naturaleza.')
qa('¿Qué es exactamente el análisis espejo?',
   'Comparar los flujos que México reporta exportar con los que los socios reportan importar desde México, y comparar la composición por etapa de exportaciones e importaciones. Si se exporta bruto y se importa procesado del mismo mineral, la transformación se hace afuera.')
qa('¿El CCV no duplica al Ghosh?',
   'No: lo complementa. El Ghosh es discreto (dos o tres cortes de MIP) y el CCV es una serie anual continua de 34 años; miden lo mismo (encadenamiento hacia adelante) desde ángulos distintos.')
qa('¿Por qué el HHI y no otra medida de concentración (C4, entropía)?',
   'El HHI usa toda la distribución y pondera por tamaño al cuadrado, lo que lo hace estándar y comparable con umbrales de política de competencia; además permite el “número efectivo de competidores”.')

H2('Errores a evitar')
bullet('No presentar Leontief/Ghosh como un modelo predictivo o causal: son descriptores de posición.')
bullet('Tener presente la cautela del Ghosh (arrastre ≠ captura de valor) por si preguntan por un valor alto.')

# ==================================================================
# DIAPOSITIVA 16 — COBERTURA
# ==================================================================
H1('Diapositiva 16 — Cobertura del periodo (¿están cubiertos los cálculos?)')

H2('Qué muestra')
P('Una tabla que responde, indicador por indicador, si el cálculo cubre el periodo 1992-2025, con su estado y una nota; y qué queda como opcional o pendiente. El mensaje es que sí está cubierto, y que los huecos se declaran, no se imputan.')

H2('El marco a fondo')
H3('La cobertura, indicador por indicador')
bullet('CCV (captura de valor): 1992-2025, completo (serie anual mineral-año). Precios USGS: 1990-2025.')
bullet('HHI (concentración): 1994-2024. 1994-2003 reconstruido por régimen con el USGS Minerals Yearbook (manganeso, fluorita, grafito, cobre); 2004-2020 aproximado por régimen; 2021-2024 por mina. Hueco declarado: 1992-1993 (privatización en curso, sin estructura por empresa).')
bullet('Comercio por etapa: 1992-2024, con detalle por fracción arancelaria.')
bullet('Encadenamientos MIP (Leontief/Ghosh): tres cortes, 2008 · 2013 · 2018 (son “fotos” por naturaleza; 2008 como referencia histórica).')
bullet('Comparación internacional (Ghosh sector, OECD ICIO): mismos tres cortes 2008 · 2013 · 2018.')
bullet('DVA / reprocesamiento (OECD ICIO): serie anual 1995-2020.')
bullet('Opcional o pendiente: actualización RAS 2020/2023 (inviable por mineral, porque la MIP a nivel Clase solo existe para años base); encadenamiento por estado con cocientes de localización (requiere la matriz de PIB estatal por sector de INEGI); comercio bilateral por socio.')
H3('El principio que la sostiene: declarar, no imputar')
P('El criterio metodológico es no rellenar los huecos con supuestos. Donde el dato no existe —el HHI de 1992-1993, el RAS por mineral—, se declara explícitamente y se explica por qué. Esto vuelve el trabajo auditable y honesto, y es preferible a inventar una serie “completa” falsa. Es un rasgo defendible, no una carencia.')
P('Detalle técnico útil por si preguntan por el HHI antiguo o el regional: la reconstrucción 1994-2003 se hace “por régimen” (a partir del líder y los grupos documentados en la Tabla 2 del USGS, tratando el residual como atomístico: es una cota inferior). El “Ghosh por estado” pendiente exigiría regionalizar la matriz nacional con métodos de no-encuesta (cocientes de localización, SLQ/FLQ), que requieren una matriz sector×estado que INEGI solo expone de forma interactiva.')

H2('Cómo exponerla')
guion('Sí, el periodo está cubierto: el CCV y el comercio son series continuas, el HHI llega de 1994 a 2024, y los encadenamientos tienen tres cortes, complementados con la comparación internacional y la descomposición de valor agregado. Donde el dato no existe —el HHI de 1992-93 o un RAS por mineral— lo declaro explícitamente; no lo invento.')

H2('Preguntas probables del jurado')
qa('¿Por qué el HHI no llega hasta 1992?',
   'Porque antes de 1994 no hay estructura de mercado por empresa reconstruible: 1992-1993 fue el pico de la privatización, con la propiedad en plena reconfiguración. 1994-2003 sí se reconstruyó, por régimen, con el USGS Minerals Yearbook.')
qa('¿Por qué solo tres cortes de la matriz insumo-producto?',
   'Porque a nivel de Clase (el mineral), los datos abiertos solo llegan a 2013 y 2018; 2008 se añade como referencia histórica en base distinta. Una actualización por mineral con el método RAS a 2020/2023 es inviable, porque la MIP a nivel Clase solo existe para años base.')
qa('¿El año 2025 está completo?',
   'Las series duras cierran en 2024 por el rezago de publicación; el CCV llega a 2025 donde el comercio ya está disponible. Se declara el cierre.')
qa('¿No debilita la tesis tener huecos?',
   'Al contrario: declararlos y explicar por qué existen es más riguroso que imputar datos. La descripción se sostiene en la convergencia de varios indicadores, no en una sola serie.')

H2('Errores a evitar')
bullet('No ocultar los huecos ni maquillarlos: preséntalos como decisión metodológica transparente.')
bullet('No confundir “corte” con “serie”: los encadenamientos MIP son fotos; el CCV y el comercio, series.')

# cierre
P('')
P('Para las cifras exactas de cada indicador y las fórmulas completas, remitirse al «Resumen metodológico y de resultados 2026-09-07». Esta guía se centra en el porqué y el cómo explicar.', size=10.5, italic=True, color=MUT)

doc.save(OUT)
print('escrito:', OUT)
print('parrafos:', len(doc.paragraphs))
