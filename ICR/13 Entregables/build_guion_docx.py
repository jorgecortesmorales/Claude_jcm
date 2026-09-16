# -*- coding: utf-8 -*-
"""Guion hablado de la exposición (protocolo ICR), desde la diapositiva 2.
Solo el guion, en registro oral, con transiciones y <= 2 min por diapositiva.
Asume la versión con enclave estructural. python-docx, formato tesis."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTDIR = r"C:\Users\Jorge\OneDrive\Escritorio\Claude CODE\ICR\13 Entregables\Resumenes descriptivos"
FECHA = "2026-09-08"
OUT = os.path.join(OUTDIR, f"Guion de la exposicion (desde diapo 2) {FECHA}.docx")
INK = RGBColor(0x21,0x1d,0x18); COP = RGBColor(0xb0,0x57,0x1e); MUT = RGBColor(0x6b,0x64,0x5a)

doc = Document()
st = doc.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.3
sec=doc.sections[0]; sec.page_width=Cm(21.59); sec.page_height=Cm(27.94)
for m in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(sec,m,Cm(3))
h2=doc.styles['Heading 2']; h2.font.name='Times New Roman'

def slide(num, titulo, mins, parrafos):
    p=doc.add_heading(level=2); p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(f'Diapositiva {num} · {titulo}'); r.font.size=Pt(13); r.bold=True; r.font.color.rgb=COP
    pt=doc.add_paragraph(); pt.paragraph_format.space_after=Pt(6)
    rt=pt.add_run(f'(≈ {mins})'); rt.font.size=Pt(10.5); rt.italic=True; rt.font.color.rgb=MUT
    for txt in parrafos:
        pp=doc.add_paragraph(); pp.paragraph_format.space_after=Pt(8); pp.paragraph_format.line_spacing=1.3
        rr=pp.add_run(txt); rr.font.size=Pt(12.5)

# ---------- Portada ----------
doc.add_paragraph().add_run('ICR · Maestría en Economía · UAM Azcapotzalco').font.size=Pt(10.5)
doc.paragraphs[-1].runs[0].font.color.rgb=COP
p=doc.add_paragraph(); r=p.add_run('Guion de la exposición'); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=INK
p=doc.add_paragraph(); r=p.add_run('Lo que se dice en voz alta, diapositiva por diapositiva (desde la 2)'); r.italic=True; r.font.size=Pt(13); r.font.color.rgb=MUT
pr=doc.add_paragraph(); pPr=pr._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); bot=OxmlElement('w:bottom')
bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'12'); bot.set(qn('w:space'),'6'); bot.set(qn('w:color'),'b0571e')
pbdr.append(bot); pPr.append(pbdr)
n=doc.add_paragraph(); rn=n.add_run('Este guion está pensado para leerse/decirse con naturalidad, no palabra por palabra. Cada diapositiva no pasa de dos minutos. La diapositiva 7 (marco teórico) está escrita para la versión con enclave estructural; si expones la versión fiel al protocolo, cambia solo esa diapositiva por el estructuralismo de Prebisch más los cuatro marcos analíticos (economía industrial, integración vertical, institucional y cadenas de valor globales).')
rn.font.size=Pt(11); rn.italic=True; rn.font.color.rgb=MUT

# =========================== GUION ===========================

slide(2, 'Planteamiento del problema', '1 min 50 s', [
 'Para entender los mercados de los minerales críticos en México hay que volver a los años noventa. Entre 1988 y 1994 el sector vivió una transformación de fondo: se desincorporaron cerca de seis punto seis millones de hectáreas de reservas y se privatizaron las grandes paraestatales —Cananea en 1990, Autlán en 1993—.',
 'La pieza que ordenó todo fue la Ley Minera de 1992, que estableció concesiones de hasta cincuenta años renovables, sin condicionarlas a que se generara valor agregado dentro del país. Esa decisión configuró, de manera durable, la estructura que tenemos hoy: unos pocos consorcios de capital mayoritariamente nacional —Grupo México, Peñoles, Frisco— con posiciones dominantes. Lo documentan Sariego y Delgado Wise.',
 'El resultado, tres décadas después, es paradójico: México es primer productor mundial de plata y actor central en varios minerales, pero importa buena parte de los bienes de alta tecnología que usan esos mismos minerales. La reforma a la Ley Minera de 2023 es el telón de fondo reciente.',
 'Quiero distinguir dos cosas. El problema factual es un hecho: se exporta el mineral en bruto y no se desarrolló la transformación aguas abajo. El problema de investigación es otro: describir con rigor, mineral por mineral, cómo son esos mercados y en qué eslabón de la cadena participa el país. Esa descripción, hoy, no existe de forma sistemática.',
])

slide(3, 'Justificación: tres literaturas que no se cruzan', '2 min', [
 'El tema se ha estudiado, pero desde tres frentes que casi no dialogan entre sí. El primero es la economía política de la privatización: explica cómo se privatizó y se concentró el sector —Sariego, Delgado Wise— pero no mide la cadena de valor. El segundo son los estudios ambientales y sociales del extractivismo, como Azamar y Ponce: describen muy bien el impacto territorial del modelo primario-exportador, pero con poca medición económica desagregada. El tercero son los estudios macroeconómicos, como el de Aliphat y colegas: cuantifican el aporte del sector al PIB y a las exportaciones, pero se quedan en el agregado nacional.',
 'La brecha, entonces, no es que falten estudios, sino que ninguno construye —de forma sistemática y por mineral— los indicadores que unen estructura de mercado, cadena de valor e inserción global. Ese acervo empírico es, justamente, mi aporte.',
 'Teóricamente me anclo en la tradición estructuralista latinoamericana, en Prebisch. Su tesis, la de Prebisch-Singer, sostiene que los precios de las materias primas se deterioran de forma secular frente a los de las manufacturas: el exportador primario tiene que entregar cada vez más para comprar lo mismo de bienes industriales. Es el marco de fondo de por qué importa dónde nos quedamos en la cadena.',
 'Y es una investigación que corresponde de lleno al campo de Empresas, Finanzas e Innovación, porque caracteriza a las empresas del sector, su estructura de industria y la organización de sus cadenas.',
])

slide(4, 'Preguntas de investigación', '1 min', [
 'De ahí se desprende la pregunta central: cómo son los mercados de los diez minerales críticos en México —en su estructura extractiva y en su articulación con la industria de transformación— y cómo se inserta el país en sus cadenas de valor globales.',
 'Y tres preguntas subsidiarias, una por cada fase del trabajo. La primera, por la estructura del mercado extractivo de cada mineral —los actores, la concentración medida con el HHI— y su encadenamiento hacia adelante. La segunda, por la cadena de valor local: ¿existen empresas que usen estos minerales como insumo dentro del país? Y la tercera, por la inserción global: cómo y con qué magnitud participa México en las cadenas mundiales de estos minerales.',
])

slide(5, 'Objetivos', '55 s', [
 'Los objetivos son el reflejo de esas preguntas. El general: caracterizar los mercados y las cadenas de valor —locales y globales— de los minerales críticos, en sus eslabones extractivo e industrial.',
 'Y tres específicos, uno a uno con las preguntas: describir la estructura extractiva por mineral y su encadenamiento hacia adelante, con Leontief, Ghosh y captura de valor; mapear la cadena local y la industria transformadora doméstica; y situar a México en las cadenas globales, por eslabón y magnitud, con los datos de comercio.',
])

slide(6, 'Hipótesis', '1 min 15 s', [
 'La hipótesis general nombra el rasgo que espero describir: una estructura extractiva concentrada y una articulación débil con la industria de transformación doméstica. Es lo que llamo enclave estructural.',
 'De ahí, tres hipótesis subsidiarias, y quiero subrayar algo importante: las enuncio como co-ocurrencias descriptivas, no como relaciones causales. La primera: donde hay más concentración en la extracción tiende a coincidir un menor encadenamiento hacia adelante. La segunda: en varios minerales la transformación doméstica es incipiente o inexistente, se exporta en bruto y se importa el bien procesado. Y la tercera: México se inserta en las cadenas globales en los eslabones primarios, con escasa captura de valor.',
])

slide(7, 'Marco teórico: el enclave estructural', '2 min', [
 'El concepto que ordena todo el trabajo es el enclave estructural, y llego a él por contraste. En la teoría de la dependencia, Cardoso y Faletto definen la economía de enclave por la propiedad extranjera del capital: el sector exportador está en manos de afuera y la renta sale del país con la empresa.',
 'Pero el caso mexicano no encaja ahí, porque el capital es mayoritariamente nacional. Y sin embargo se comporta como enclave: se extrae y se exporta en bruto sin desarrollar la cadena industrial local. Por eso hablo de enclave estructural. El criterio se desplaza: ya no es de quién es la mina, sino dónde se detiene la cadena. Me apoyo en el neo-extractivismo de Svampa, que describe cómo la región profundizó su perfil primario-exportador durante el auge de las materias primas, y en la raíz estructuralista de Prebisch.',
 'Para volver esto medible uso las herramientas de la economía industrial y de las cadenas de valor: la estructura de mercado, la integración vertical de las firmas y la posición en la cadena global. Con ellas, tres descriptores —el HHI, el índice de Ghosh y el coeficiente de captura de valor— hacen visible dónde la cadena existe, dónde se detiene y dónde está ausente.',
 'Un matiz que cuidaré toda la exposición: un encadenamiento estadístico alto no equivale a una cadena desarrollada. Por eso nunca leo un solo indicador; el enclave estructural aparece en la lectura conjunta de varios.',
])

slide(8, 'Enfoque y horizonte temporal', '1 min 50 s', [
 'El enfoque es descriptivo y de métodos mixtos. Descriptivo, y lo digo con claridad: no busco estimar un efecto causal, sino caracterizar con rigor los mercados. La razón es doble. Por un lado, la información encadenable por mineral son apenas dos o tres cortes de la matriz insumo-producto, insuficientes para una identificación causal creíble. Por otro, el vacío en la literatura es precisamente descriptivo: no existen estos indicadores por mineral. Un modelo causal quedaría como agenda futura.',
 'Métodos mixtos porque combino lo cuantitativo documental —la matriz insumo-producto, el comercio, los precios— con lo cualitativo, la lectura de reportes de empresas. Cada componente ve lo que el otro no: el dato sectorial no capta la integración dentro de las firmas, y el reporte corporativo sí.',
 'La unidad de observación es el mineral-año: cada observación es un mineral en un año, lo que me deja comparar entre minerales y seguir a cada uno en el tiempo. El horizonte es de 1992, la Ley Minera, a 2025, el cierre de la información. Y aviso desde ya algo que retomaré: el detalle temporal depende de la fuente. Unos indicadores son series anuales continuas; otros, cortes puntuales de la matriz. Trabajo con la granularidad que cada fuente permite, y lo declaro.',
])

slide(9, 'Fuentes de información', '1 min', [
 'Las fuentes son públicas y verificables. Para los encadenamientos, la matriz insumo-producto del INEGI. Para el comercio por etapa, UN Comtrade y DataMéxico. Para la estructura de la industria y los precios, el USGS —el Minerals Yearbook y los Mineral Commodity Summaries— y la Cámara Minera de México. Para la cadena de valor local, los reportes corporativos. Y para la comparación internacional y el valor agregado, las tablas insumo-producto inter-país de la OCDE. El marco institucional lo dan la Ley Minera y su reforma.',
])

slide(10, 'Técnicas de análisis', '2 min', [
 'A cada objetivo le corresponde una técnica, y las presento de lo simple a lo elaborado.',
 'Para la estructura extractiva, el índice de Herfindahl-Hirschman: la suma de los cuadrados de las participaciones. Como eleva al cuadrado, un líder dominante pesa mucho más que varios productores pequeños, así que distingue bien un monopolio de un mercado repartido.',
 'Para el encadenamiento, dos modelos duales de la matriz insumo-producto. El de Leontief mide hacia atrás —cuánto arrastra un mineral a sus proveedores— y en minería suele ser bajo. El de Ghosh mide hacia adelante —cuánto alimenta a las industrias que lo usan como insumo—, que es justo la pregunta de la tesis. A esos dos sumo el coeficiente de captura de valor, que da la versión anual y continua del encadenamiento hacia adelante, año con año.',
 'Para la cadena local, la demanda intermedia: qué sectores compran cada mineral dentro del país, complementada con un mapa de empresas, porque buena parte de la cadena ocurre dentro de los propios grupos.',
 'Y para la inserción global, el comercio por etapa: qué parte se exporta en bruto y qué parte procesada, con el análisis espejo —si exportamos en bruto e importamos el producto procesado del mismo mineral, la transformación se hace afuera—.',
 'Insisto en un punto: todos son descriptores, no un modelo causal. Y el índice de Ghosh se lee con cautela, como posición estructural, nunca como prueba por sí solo de que existe una cadena desarrollada.',
])

slide(11, 'Ajustes metodológicos (transición)', '40 s', [
 'Hasta aquí, el diseño. Ahora, al ejecutar el trabajo, algunas decisiones se ajustaron a lo que los datos hicieron posible. Las presento con transparencia, cada una con su porqué, porque creo que esos ajustes fortalecen la investigación.',
])

slide(12, 'Ajuste 1 · de un diseño causal a uno descriptivo', '1 min 15 s', [
 'El cambio de fondo. En una primera versión pensaba en un modelo causal de panel —la concentración explicando el encadenamiento— con un estudio de evento sobre la reforma. Lo reencuadré hacia una descripción de las cadenas de valor: el HHI, el Ghosh y la captura de valor pasan a ser descriptores, no variables de un modelo, y el aporte es construir los datos e indicadores por mineral.',
 '¿Por qué? Porque con tan pocos cortes de matriz no hay base para una identificación causal creíble, y porque la contribución realmente novedosa —ausente en la literatura— es medir y describir la desconexión de la cadena. Fue un acuerdo con el asesor y se afinó en el análisis de consistencia.',
])

slide(13, 'Ajuste 2 · qué cortes de la matriz se pueden usar', '1 min', [
 'El segundo ajuste es sobre los cortes de la matriz insumo-producto. Planteaba varios cortes actualizables. En la práctica, a nivel de Clase —el nivel del mineral— los datos abiertos solo llegan a 2013 y 2018, y añado 2008 como referencia histórica, en base distinta. Los años intermedios solo existen a un nivel más agregado, que no separa el mineral. Aun así, logro profundidad temporal: 2008, 2013 y 2018.',
])

slide(14, 'Ajuste 3 · el CCV como serie anual', '1 min', [
 'El tercer ajuste da continuidad temporal al encadenamiento hacia adelante. El índice de Ghosh es una foto de dos o tres años; el periodo que estudio abarca treinta y cuatro. Por eso construí el coeficiente de captura de valor como serie anual de 1992 a 2025: el valor unitario de lo que se exporta en bruto sobre el precio del producto de referencia. Es una segunda medida del encadenamiento hacia adelante, ahora continua.',
])

slide(15, 'Ajuste 4 · la cobertura real de las series', '1 min', [
 'El cuarto ajuste es de honestidad sobre hasta dónde llega cada dato. Algunas series no cubren toda la ventana con la misma granularidad, y eso es un límite de la fuente, no una decisión de diseño. Donde falta el dato, lo declaro. Es lo que detallo en la siguiente diapositiva.',
])

slide(16, 'Cobertura del periodo: ¿están cubiertos los cálculos?', '2 min', [
 'La respuesta corta es sí, y quiero mostrarlo con transparencia. El coeficiente de captura de valor cubre todo el periodo, de 1992 a 2025, año con año, igual que los precios de referencia. El comercio por etapa, de 1992 a 2024. El HHI, de 1994 a 2024: los años más antiguos, de 1994 a 2003, los reconstruí por régimen a partir de la estructura del USGS para los minerales donde el patrón es inequívoco —manganeso, fluorita, grafito y cobre—.',
 'Los encadenamientos de la matriz son, por naturaleza, fotos: tengo tres cortes, 2008, 2013 y 2018. Y sobre esa misma base comparo internacionalmente a México con Chile, Australia, Finlandia y Suecia, en los mismos tres cortes, y añado una descomposición de valor agregado, de 1995 a 2020.',
 'Hay huecos, y los declaro en voz alta en vez de rellenarlos con supuestos: el HHI de 1992 y 1993 no es reconstruible, porque en plena privatización no hay una estructura de mercado estable por empresa. Y quedan como opcionales o pendientes una actualización por el método RAS —inviable por mineral, porque la matriz a ese nivel solo existe para años base— y el encadenamiento por estado, que requiere una matriz de PIB estatal por sector.',
 'Este criterio —declarar, no imputar— es deliberado: prefiero una descripción honesta y auditable, sostenida en la convergencia de varios indicadores, a una serie aparentemente completa pero fabricada.',
])

slide(17, 'Cronograma', '45 s', [
 'El cronograma organiza el trabajo en doce meses. Al día de hoy, la fase de datos está completa: los indicadores construidos y validados. El frente activo es la inserción global, la síntesis y la redacción, con revisiones continuas con el asesor.',
])

slide(18, 'Contribución (cierre)', '1 min 40 s', [
 'Cierro con el mensaje central. El valor de esta tesis no está en un modelo econométrico, sino en construir y describir algo que no existía: datos e indicadores por mineral —el HHI de 1994 a 2024, los encadenamientos en tres cortes, el coeficiente de captura de valor de treinta y cuatro años, y el comercio por etapa—.',
 'A eso sumo dos piezas que reencuadran el debate. La comparación internacional muestra que ni Chile ni Australia son casos de éxito en el encadenamiento hacia adelante; el verdadero referente de integración es el modelo nórdico. Y la descomposición de valor agregado mide el enclave en dinero: cuánto del valor de la minería se queda transformado en casa y cuánto se va en crudo para procesarse afuera.',
 'Todo ordenado por un solo concepto —el enclave estructural— y con bases concretas para pensar una política industrial, diferenciadas por mineral y por eslabón ausente. Gracias.',
])

slide(19, 'Referencias', '20 s', [
 'Estas son las referencias principales del trabajo, organizadas por marco teórico, insumo-producto y método, minería en México, y fuentes de datos. La bibliografía completa está en el documento. Quedo atento a sus comentarios.',
])

doc.save(OUT)
print('escrito:', OUT)
print('parrafos:', len(doc.paragraphs))
